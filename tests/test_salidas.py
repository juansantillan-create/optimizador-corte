"""Informe, dibujos y exportación (fase F2)."""

import io
import re

import openpyxl

from optimizador import dibujo, exportar, informe
from optimizador.api import optimizar


def _resultado():
    ent = {"proyecto": {"codigo": "T2", "descripcion": "PRJ T2"},
           "piezas": [{"item": "1", "material": "FL5IN", "cantidad": 7, "ancho": 1150, "largo": 2400},
                      {"item": "2", "material": "FL5IN", "cantidad": 3, "ancho": 790, "largo": 1590}],
           "opciones": {"esfuerzo": "rapido"}, "fecha": "2026-10-20 23:20:19"}
    res = optimizar(ent)
    assert res["verificado"]
    return res


def _como_413(res):
    """El mismo resultado con los números del caso 413 de la captura (para ver el formato)."""
    m = res["materiales"][0]
    m.update({"piezas": 252, "m2_piezas": 564.583068, "layouts": 13, "planchas": 82,
              "m2_planchas": 733.9, "aprovechamiento": 76.93, "perdida": 23.07, "perimetro_m": 1630.22,
              "planchas_usadas": [{"cantidad": 82, "ancho": 3580, "largo": 2500, "tipo": "Plancha",
                                   "hoja_ancho": 3600, "hoja_largo": 2500, "supuesta": False}]})
    return res


def test_informe_con_el_formato_de_corte_certo():
    res = _como_413(_resultado())
    html = informe.html_resultado(res)
    texto = re.sub(r"<[^>]+>", " ", html)
    for esperado in ("Informe de cálculo", "20/10/26", "23:20:19", "Proyectos calculados",
                     "Lista de Materiales:", "564.583068", "733.9", "76.93%", "3580.0", "2500.0",
                     "Plancha", "Perímetro - Piezas: 1630.22 m", "Cálculo especial de perímetro: 0.00 m",
                     "Elija el material para visualizar:", "Cdad. de montajes:", "Total de chapas:",
                     "Aprovechamiento:", "Pérdida:", "23.07 %"):
        assert esperado in texto, esperado
    assert "@page { size: A4 landscape" in html
    assert "http" not in html.replace("http://www.w3.org/2000/svg", "")  # nada de internet


def test_informe_con_coma():
    res = _como_413(_resultado())
    texto = re.sub(r"<[^>]+>", " ", informe.html_resultado(res, coma=True))
    assert "564,583068" in texto and "76,93%" in texto


def test_miniatura_apilada():
    res = _resultado()
    m = max(res["materiales"][0]["montajes"], key=lambda m: m["repeticiones"])
    svg = dibujo.miniatura(m)
    assert svg.startswith("<svg")
    # una capa por repetición de más (hasta 12) + la plancha + las piezas
    capas = min(m["repeticiones"] - 1, 12)
    assert svg.count('fill="#ffffff"') == capas
    if m["repeticiones"] > 1:
        assert f"×{m['repeticiones']}" in svg
    assert svg.count('fill="#800000"') == sum(1 for p in m["piezas"] if p["tipo"] == 1)


def test_detalle_tiene_medidas_y_cortes():
    res = _resultado()
    m = res["materiales"][0]["montajes"][0]
    svg = dibujo.detalle(m)
    assert "1150 × 2400" in svg or "2400 × 1150" in svg
    assert svg.count('class="corte"') == len(m["cortes"])


def test_exportar_csv_y_xlsx():
    res = _resultado()
    csv = exportar.a_csv(res)
    assert csv.startswith("\ufeffProyecto;Material;")
    filas = csv.strip().splitlines()
    assert len(filas) == 2 and ";FL5IN;" in filas[1]
    datos = exportar.a_xlsx(res)
    libro = openpyxl.load_workbook(io.BytesIO(datos))
    hoja = libro["Consumo"]
    assert hoja.cell(1, 2).value == "Material" and hoja.cell(2, 2).value == "FL5IN"
    assert hoja.cell(2, 9).value == res["materiales"][0]["planchas"]
